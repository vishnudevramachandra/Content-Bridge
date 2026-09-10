"""Shared minimalist theme constants for the ContentBridge review deck.

Kept in one place so every slide-building function (added incrementally in
build_deck.py) stays visually consistent without repeating color/font
literals everywhere.
"""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# --- Palette -----------------------------------------------------------
BACKGROUND = RGBColor(0xFF, 0xFF, 0xFF)   # white
TEXT_DARK = RGBColor(0x22, 0x29, 0x2F)    # near-black charcoal, body/title text
ACCENT = RGBColor(0x0F, 0x6E, 0x8C)       # teal-blue, titles/highlights/diagram accents
MUTED = RGBColor(0x8A, 0x94, 0xA6)        # muted gray, subtitles/footers
ACCENT_LIGHT = RGBColor(0xD9, 0xEE, 0xF2)  # pale teal tint, for boxes/backgrounds behind diagrams
WARN = RGBColor(0xB1, 0x42, 0x2C)         # muted brick red, used sparingly for "flagged/mismatch" states

# --- Fonts ---------------------------------------------------------------
FONT = "Calibri"

# --- Slide geometry (16:9 widescreen) ------------------------------------
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Consistent margins used by content slides (title/body layout)
MARGIN_X = Inches(0.7)
TITLE_TOP = Inches(0.4)
TITLE_H = Inches(0.9)
BODY_TOP = Inches(1.5)


def set_background(slide):
    """Plain white background for every slide, applied explicitly rather
    than relying on the default theme (which could differ per template)."""
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = BACKGROUND


def add_slide_title(slide, text):
    """Consistent top-left title + thin accent underline, used by every
    content slide (2 onward) so headers don't drift slide to slide."""
    from pptx.enum.text import PP_ALIGN

    title_box = slide.shapes.add_textbox(MARGIN_X, TITLE_TOP, Inches(11.9), TITLE_H)
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.name = FONT
    run.font.color.rgb = TEXT_DARK

    line = slide.shapes.add_shape(1, MARGIN_X, Inches(1.25), Inches(1.2), Pt(2.5))
    line.fill.solid()
    line.fill.fore_color.rgb = ACCENT
    line.line.fill.background()


def add_box(slide, x, y, w, h, text, fill=None, text_color=None, font_size=14,
            bold=False, outline=True):
    """Rounded rectangle with centered text — the basic unit used for every
    diagram box across slides 3-6 (system panels, ERD-style entities)."""
    fill = fill if fill is not None else ACCENT_LIGHT
    text_color = text_color if text_color is not None else TEXT_DARK

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if outline:
        shape.line.color.rgb = ACCENT
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False

    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = text_color
    return shape


def add_card(slide, x, y, w, h, title, lines, title_color=None):
    """Light rounded-rectangle 'card' with a bold left-aligned title and a
    few bulleted lines beneath it — used for content previews (posts,
    pages) where a plain labeled box isn't enough."""
    title_color = title_color if title_color is not None else ACCENT

    add_box(slide, x, y, w, h, "", outline=True)  # background only

    box = slide.shapes.add_textbox(x + Inches(0.15), y + Inches(0.08),
                                    w - Inches(0.3), h - Inches(0.16))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = title
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.name = FONT
    run.font.color.rgb = title_color

    for line in lines:
        p = tf.add_paragraph()
        p.space_before = Pt(2)
        run = p.add_run()
        run.text = f"\u2022 {line}"
        run.font.size = Pt(11.5)
        run.font.name = FONT
        run.font.color.rgb = TEXT_DARK


def add_content_card(slide, x, y, w, h, kind, title, excerpt):
    """Single WordPress post/page preview card: type label, title, and a
    short excerpt \u2014 used when each item should read as its own piece of
    content rather than a line in a shared list."""
    add_box(slide, x, y, w, h, "", outline=True)  # background only

    box = slide.shapes.add_textbox(x + Inches(0.15), y + Inches(0.08),
                                    w - Inches(0.3), h - Inches(0.16))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = kind.upper()
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.name = FONT
    run.font.color.rgb = MUTED

    p = tf.add_paragraph()
    p.space_before = Pt(2)
    run = p.add_run()
    run.text = title
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.name = FONT
    run.font.color.rgb = TEXT_DARK

    p = tf.add_paragraph()
    p.space_before = Pt(3)
    run = p.add_run()
    run.text = excerpt
    run.font.size = Pt(10)
    run.font.italic = True
    run.font.name = FONT
    run.font.color.rgb = MUTED


def add_field_tag(slide, x, y, w, h, field_name, value):
    """Small highlighted 'field: value' pill \u2014 used to call out the one
    custom field that's actually driving a mapping decision on a slide."""
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = ACCENT_LIGHT
    box.line.color.rgb = ACCENT
    box.line.width = Pt(1.25)
    box.shadow.inherit = False

    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = field_name
    run.font.size = Pt(11)
    run.font.name = FONT
    run.font.color.rgb = MUTED

    p = tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = value
    run.font.size = Pt(18)
    run.font.bold = True
    run.font.name = FONT
    run.font.color.rgb = ACCENT


def add_mini_table(slide, x, y, w, h, headers, rows, col_widths=None, highlight_rows=None):
    """Small header+rows grid, used for the Strapi entity previews (and
    reused for the mapping-rules table on the end-product slide).

    highlight_rows: 0-based indices (into `rows`) to render as the "found
    match" \u2014 bold text on a pale accent fill \u2014 versus plain rows.
    """
    highlight_rows = highlight_rows or set()
    n_rows = len(rows) + 1
    n_cols = len(headers)
    graphic_frame = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    table = graphic_frame.table

    if col_widths:
        total = sum(col_widths)
        for c, frac in enumerate(col_widths):
            table.columns[c].width = int(w * frac / total)

    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = str(htext)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
        cell.margin_top = Pt(2)
        cell.margin_bottom = Pt(2)
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = BACKGROUND
                run.font.name = FONT

    for ri, row in enumerate(rows, start=1):
        is_hit = (ri - 1) in highlight_rows
        for c, val in enumerate(row):
            cell = table.cell(ri, c)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT_LIGHT if is_hit else BACKGROUND
            cell.margin_top = Pt(1)
            cell.margin_bottom = Pt(1)
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.size = Pt(10.5)
                    run.font.bold = is_hit
                    run.font.color.rgb = ACCENT if is_hit else MUTED
                    run.font.name = FONT
    return graphic_frame


def add_connector(slide, x1, y1, x2, y2, dashed=False, color=None):
    """Straight line between two diagram boxes (ERD-style relation lines)."""
    color = color if color is not None else MUTED
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(1.25)
    if dashed:
        conn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return conn
